from __future__ import annotations
from io import BytesIO
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def build_excel_report(kpi_dict, branch_df, detail_df, incidents_df, schedule_df):
    output = BytesIO()
    summary = pd.DataFrame({
        "Indicador": ["Faltantes", "Merma", "Exactitud inventario", "A cobro", "Incidencias abiertas", "Registros analizados"],
        "Valor": [kpi_dict["shortage"], kpi_dict["shrinkage"], kpi_dict["accuracy"], kpi_dict["chargeable"], kpi_dict["open_incidents"], kpi_dict["records"]],
    })
    discrepancies = detail_df[detail_df["difference_qty"] != 0].copy()
    shrinkage = detail_df[detail_df["damaged_qty"] > 0].copy()
    chargeable = detail_df[detail_df["chargeable_amount"] > 0].copy()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Executive Summary", index=False)
        branch_df.to_excel(writer, sheet_name="Branch Performance", index=False)
        detail_df.to_excel(writer, sheet_name="Inventory Detail", index=False)
        discrepancies.to_excel(writer, sheet_name="Discrepancies", index=False)
        shrinkage.to_excel(writer, sheet_name="Shrinkage", index=False)
        chargeable.to_excel(writer, sheet_name="Chargeable", index=False)
        incidents_df.to_excel(writer, sheet_name="Incidents", index=False)
        schedule_df.to_excel(writer, sheet_name="Visit Schedule", index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center")
            for col_idx, col_cells in enumerate(ws.columns, 1):
                max_len = max((len(str(c.value)) if c.value is not None else 0) for c in list(col_cells)[:100])
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 11), 34)
        ws = writer.book["Executive Summary"]
        ws["B2"].number_format = '$#,##0.00'
        ws["B3"].number_format = '$#,##0.00'
        ws["B4"].number_format = '0.00%'
        ws["B5"].number_format = '$#,##0.00'
    output.seek(0)
    return output
